"""Deck skill (scripts/deck/make-deck.py): the parts that run without Marp or a browser. Created 2026-10-25."""
import importlib.util
import os
import subprocess
import sys
import zipfile

import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIX = os.path.join(ROOT, "tests", "fixtures", "deck")
_spec = importlib.util.spec_from_file_location("make_deck", os.path.join(ROOT, "scripts", "deck", "make-deck.py"))
deck = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(deck)


def _transcript():
    return open(os.path.join(FIX, "lecture-social-fact.ru.txt"), encoding="utf-8").read()


def test_transcript_draft_has_title_and_no_empty_slide():
    title, slides = deck.slides_from_transcript(_transcript())
    assert title.startswith("Введение в социологию")
    md = deck.slides_to_markdown(title, slides, "default", "ru")
    parsed = deck.marp_slides(md)
    assert len(parsed) == len(slides)          # front matter is not a slide of its own
    assert parsed[0]["lead"] and all(p["heading"] for p in parsed)


def test_cut_bullet_keeps_full_sentence_in_notes():
    long = "Это предложение " + "очень " * 20 + "длинное."
    s = deck.slide_from_sentences(["Заголовок.", long])
    assert s["bullets"][0].endswith("…") and len(s["bullets"][0]) <= deck.BULLET_MAX + 1
    assert long in s["notes"]


def test_question_mark_survives():
    assert deck.strip_terminal("Why does it continue?") == "Why does it continue?"
    assert deck.strip_terminal("A point.") == "A point"


def test_language_is_the_majority_script():
    assert deck.mostly_cyrillic(_transcript())
    assert not deck.mostly_cyrillic('Culture: «Бьёт — значит любит» means "he beats you means he loves you", '
                                    "and nobody believes it literally, but it lowers the cost of silence.")


def test_pptx_zip_reader_on_fixture():
    raw = deck.pptx_slides_zip(os.path.join(FIX, "sample.pptx"))
    assert raw[0][0] == "Социология города"
    assert any(n for _, _, n in raw)            # the notes slide is read


def test_marp_slides_levels_notes_directives():
    md = ("---\nmarp: true\n---\n\n<!-- _class: lead -->\n# T\n\nSub\n\n---\n\n## H\n\n- a **b**\n  - c\n\n"
          "![w](wheel.png)\n\n<!--\nsay this\n-->\n")
    s = deck.marp_slides(md)
    assert s[0]["lead"] and s[0]["items"] == [(0, "Sub")] and s[0]["notes"] == []
    assert s[1]["items"] == [(0, "a b"), (1, "c")]
    assert "say this" in s[1]["notes"] and any(n.startswith("[image:") for n in s[1]["notes"])


def test_write_pptx_is_editable(tmp_path):
    pytest.importorskip("pptx")
    from pptx import Presentation
    md = deck.slides_to_markdown(*deck.slides_from_transcript(_transcript()), "default", "ru")
    out = tmp_path / "d.pptx"
    assert deck.write_pptx(md, str(out))
    prs = Presentation(str(out))
    assert prs.slides[0].shapes.title.text.startswith("Введение")
    assert prs.slides[1].has_notes_slide
    with zipfile.ZipFile(out) as z:             # text, not pictures
        assert not [n for n in z.namelist() if n.startswith("ppt/media/")]


def test_draft_only_cli(tmp_path):
    r = subprocess.run([sys.executable, os.path.join(ROOT, "scripts", "deck", "make-deck.py"),
                        os.path.join(FIX, "lecture-social-fact.ru.txt"), "--out", str(tmp_path), "--draft-only"],
                       capture_output=True, text=True, stdin=subprocess.DEVNULL)
    assert r.returncode == 0, r.stderr
    d = next(tmp_path.iterdir())
    assert (d / "deck.md").exists() and (d / "notes.md").exists() and (d / "manifest.json").exists()
