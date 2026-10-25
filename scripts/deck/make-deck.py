#!/usr/bin/env python3
"""make-deck.py: a lecture deck from spoken or typed words, or from an old PPTX, written as Marp markdown
and rendered to an editable PPTX and a PDF on the box.

Created 2026-10-25 (Deck v1; first case a teaching professor's lectures). Stdlib Python, python-pptx for the
editable PPTX (the deck venv, see docs/DECK.md), the Marp CLI (npx) and a headless chromium for PDF and previews.

    make-deck.py <input> --out <dir> [--lang auto|ru|en|sk] [--theme default|gaia|uncover] [--title ...]
                 [--draft-only] [--slug ...]

Inputs
    .md    Marp markdown (front matter with "marp: true"): passed through as is (this is how a deck.md the
           reasoner refined is rendered).
           Plain markdown with "#" headings: one slide per heading, first heading is the title slide.
    .txt   a transcript (speech) or typed notes: structured into slides by the fixed rule set below, no model call.
    audio  .ogg .oga .opus .mp3 .m4a .wav .webm: transcribed first (CC_STT_URL, the OpenAI transcription shape,
           else ElevenLabs with ELEVENLABS_API_KEY), the words kept as transcript.txt, then as .txt.
    .pptx  an old deck: text per slide (python-pptx if installed, else zipfile + XML of ppt/slides/slide*.xml),
           title placeholder -> heading, other paragraphs -> bullets, notes slide -> speaker notes.

Rule set (transcript and plain text)
    1. Title slide: --title, else the first line when it is short (<= 80 chars) and not a full sentence, else the
       first sentence cut to 80 chars.
    2. Paragraphs (blank-line separated) are slides. Without paragraphs, every 5 sentences form a slide.
    3. One idea per slide: the first sentence of the paragraph becomes the heading (cut at a clause boundary,
       at most 60 chars). The next sentences become bullets, at most 6 per slide.
    4. A bullet longer than 90 chars is cut at a clause boundary and gets an ellipsis; the full sentence
       goes to the speaker notes.
    5. Leftover sentences (beyond 6 bullets) go to the speaker notes verbatim.
    6. A paragraph with one sentence only becomes a statement slide (heading only, sentence in the notes).
    The rules make a usable draft; the reasoner is expected to rewrite deck.md (--draft-only) before rendering.

Output: <out>/<slug>/deck.md, deck.pptx (editable: real title and bullet placeholders, speaker notes),
deck.pdf (Marp), preview/deck.NNN.png, notes.md, manifest.json.
--draft-only stops after deck.md + notes.md + manifest.json so a reasoner can refine deck.md, then run this
script again on that deck.md (it has front matter, so it passes through) to render.
"""
import argparse
import datetime as dt
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import time
import unicodedata
import zipfile
import zlib
import glob
import urllib.request
import uuid
import xml.etree.ElementTree as ET

THEMES = ("default", "gaia", "uncover")
LANGS = ("auto", "ru", "en", "sk")
AUDIO_EXT = (".ogg", ".oga", ".opus", ".mp3", ".m4a", ".wav", ".webm")
MAX_BULLETS = 6
BULLET_MAX = 90
HEADING_MAX = 60
TITLE_MAX = 80
SENTENCES_PER_SLIDE = 5

# Chromium for Marp: CHROME_PATH, else the newest Playwright headless shell under ~/.cache/ms-playwright.
# CHROME_LIBS: extra library dirs (LD_LIBRARY_PATH) on a box without chromium's system libraries.
CHROME_GLOB = os.path.expanduser("~/.cache/ms-playwright/chromium_headless_shell-*/chrome-headless-shell-linux*/"
                                 "chrome-headless-shell")
CHROME_LIBS_DEFAULT = ":".join(os.path.expanduser(p) for p in (
    "~/.local/chromelibs/root/usr/lib/aarch64-linux-gnu", "~/.local/chromelibs/root/lib/aarch64-linux-gnu",
    "~/.local/chromelibs/root/usr/lib/x86_64-linux-gnu", "~/.local/chromelibs/root/lib/x86_64-linux-gnu"))
MARP_DEFAULT = "npx --yes @marp-team/marp-cli"

WORDS = {
    "ru": {"notes": "Заметки докладчика", "slide": "Слайд", "title": "Лекция"},
    "en": {"notes": "Speaker notes", "slide": "Slide", "title": "Lecture"},
    "sk": {"notes": "Poznámky prednášajúceho", "slide": "Snímka", "title": "Prednáška"},
}

# ---------------------------------------------------------------- text helpers

_TRANSLIT = {
    "а": "a", "б": "b", "в": "v", "г": "g", "д": "d", "е": "e", "ё": "e", "ж": "zh", "з": "z", "и": "i",
    "й": "y", "к": "k", "л": "l", "м": "m", "н": "n", "о": "o", "п": "p", "р": "r", "с": "s", "т": "t",
    "у": "u", "ф": "f", "х": "kh", "ц": "ts", "ч": "ch", "ш": "sh", "щ": "shch", "ъ": "", "ы": "y", "ь": "",
    "э": "e", "ю": "yu", "я": "ya",
}


def slugify(text, limit=48):
    s = text.lower()
    s = "".join(_TRANSLIT.get(ch, ch) for ch in s)
    s = unicodedata.normalize("NFKD", s)
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    s = re.sub(r"[^a-z0-9]+", "-", s).strip("-")
    if len(s) > limit:
        s = s[:limit].rsplit("-", 1)[0] or s[:limit]
    return s or "deck"


_ABBREV = re.compile(r"\b(т\.е|т\.д|т\.п|др|напр|e\.g|i\.e|etc|vs|Dr|Mr|Mrs|Prof|cf|св|ст|гг|тыс|млн|млрд)\.$",
                     re.IGNORECASE)


def split_sentences(text):
    """Sentence split on . ! ? followed by whitespace and an upper-case letter, digit, quote or dash.
    Keeps common abbreviations together. Works for Cyrillic and Latin."""
    text = re.sub(r"\s+", " ", text.strip())
    if not text:
        return []
    parts = re.split(r"(?<=[.!?…])\s+(?=[\"'«„(\-–—]?[A-ZА-ЯЁÁ-Ž0-9])", text)
    out = []
    for p in parts:
        p = p.strip()
        if not p:
            continue
        if out and _ABBREV.search(out[-1]):
            out[-1] = out[-1] + " " + p
        else:
            out.append(p)
    return out


def cut_at_clause(s, limit, ellipsis):
    """Shorten s to <= limit chars, preferring a clause boundary (, ; : dash), else a word boundary."""
    s = s.strip()
    if len(s) <= limit:
        return s
    head = s[:limit]
    best = -1
    for sep in (" — ", " – ", " - ", "; ", ": ", ", "):
        i = head.rfind(sep)
        if i > limit // 3 and i > best:
            best = i
    if best < 0:
        best = head.rfind(" ")
        if best < limit // 3:
            best = limit
    cut = s[:best].rstrip(" ,;:—–-")
    return cut + ("…" if ellipsis else "")


def strip_terminal(s):
    """Drop a closing period, colon, semicolon or comma; a question or exclamation stays (it is the point)."""
    return s.rstrip(" .…:;,")


def short_bullet(s):
    """A bullet of at most BULLET_MAX chars: cut at a clause boundary, with an ellipsis when cut."""
    if len(s) <= BULLET_MAX:
        return strip_terminal(s)
    return strip_terminal(cut_at_clause(s, BULLET_MAX, False)) + "…"


def has_cyrillic(text):
    return re.search(r"[Ѐ-ӿ]", text) is not None


def mostly_cyrillic(text):
    """True when at least a third of the letters are Cyrillic (one quoted Russian proverb does not make a deck ru)."""
    letters = [ch for ch in text if ch.isalpha()]
    return bool(letters) and sum(1 for ch in letters if "Ѐ" <= ch <= "ӿ") * 3 >= len(letters)


# ---------------------------------------------------------------- slide model
# A slide is a dict: {"heading": str, "bullets": [str], "notes": [str], "title": bool, "body": str|None}
# "body" is raw markdown kept verbatim (plain-markdown input), used instead of bullets when present.


def slides_from_transcript(text, title=None, lang="ru"):
    text = text.replace("\r\n", "\n")
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    if not paragraphs:
        sys.exit("make-deck: the input has no text")

    # 1. title
    first_line = paragraphs[0].split("\n", 1)[0].strip()
    if title:
        deck_title = title.strip()
    elif len(first_line) <= TITLE_MAX and not re.search(r"[.!?]\s*$", first_line) and (
            "\n" in paragraphs[0] or len(paragraphs) > 1):
        deck_title = strip_terminal(first_line)
        rest = paragraphs[0].split("\n", 1)[1].strip() if "\n" in paragraphs[0] else ""
        paragraphs = ([rest] if rest else []) + paragraphs[1:]
    else:
        deck_title = strip_terminal(cut_at_clause(split_sentences(paragraphs[0])[0], TITLE_MAX, False))

    # 2. slide units
    units = []
    if len(paragraphs) <= 1:
        sents = split_sentences(" ".join(paragraphs))
        for i in range(0, len(sents), SENTENCES_PER_SLIDE):
            units.append(sents[i:i + SENTENCES_PER_SLIDE])
    else:
        for p in paragraphs:
            lines = [ln.strip() for ln in p.split("\n") if ln.strip()]
            # a paragraph that is already a list (-, *, •, 1.) keeps its items as bullets
            if len(lines) > 1 and all(re.match(r"^([-*•]|\d+[.)])\s+", ln) for ln in lines[1:]):
                items = [re.sub(r"^([-*•]|\d+[.)])\s+", "", ln) for ln in lines[1:]]
                units.append([lines[0]] + items)
            else:
                units.append(split_sentences(" ".join(lines)))

    slides = [{"heading": deck_title, "bullets": [], "notes": [], "title": True, "body": None}]
    for sents in units:
        if not sents:
            continue
        slides.append(slide_from_sentences(sents))
    return deck_title, slides


def slide_from_sentences(sents):
    heading = strip_terminal(cut_at_clause(sents[0], HEADING_MAX, False))
    notes = []
    if len(sents) == 1:
        return {"heading": heading, "bullets": [], "notes": [sents[0]], "title": False, "body": None}
    if len(sents[0]) > HEADING_MAX:
        notes.append(sents[0])
    bullets = []
    for s in sents[1:1 + MAX_BULLETS]:
        bullets.append(short_bullet(s))
        if len(s) > BULLET_MAX:
            notes.append(s)  # the full sentence; the bullet is cut
    notes.extend(sents[1 + MAX_BULLETS:])
    return {"heading": heading, "bullets": bullets, "notes": notes, "title": False, "body": None}


def slides_from_plain_markdown(text, title=None):
    """One slide per # or ## heading. Bullets kept; more than MAX_BULLETS move to notes; paragraphs stay."""
    text = text.replace("\r\n", "\n")
    chunks = re.split(r"(?m)^(?=#{1,2}\s)", text)
    slides = []
    preamble = ""
    for ch in chunks:
        if not ch.strip():
            continue
        m = re.match(r"^(#{1,2})\s+(.*)\n?", ch)
        if not m:
            preamble += ch
            continue
        heading = m.group(2).strip()
        body = ch[m.end():].strip("\n")
        lines = body.split("\n") if body else []
        bullets = [ln for ln in lines if re.match(r"^\s*[-*]\s+", ln)]
        notes = []
        if len(bullets) > MAX_BULLETS:
            extra = bullets[MAX_BULLETS:]
            notes = [re.sub(r"^\s*[-*]\s+", "", b) for b in extra]
            body = "\n".join(ln for ln in lines if ln not in extra)
        slides.append({"heading": heading, "bullets": [], "notes": notes, "title": False, "body": body})
    if not slides:
        return slides_from_transcript(text, title)
    slides[0]["title"] = True
    if title:
        slides[0]["heading"] = title
    deck_title = slides[0]["heading"]
    if preamble.strip():
        slides[0]["notes"].insert(0, preamble.strip())
    return deck_title, slides


# ---------------------------------------------------------------- pptx reading

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}


def _para_text(p_el):
    return "".join(t.text or "" for t in p_el.iter("{%s}t" % NS["a"])).strip()


def _shape_paragraphs(root, skip_types=("sldNum", "ftr", "dt")):
    """Yield (placeholder_type, [paragraph texts]) per shape, in document order."""
    for sp in root.iter("{%s}sp" % NS["p"]):
        ph = sp.find("p:nvSpPr/p:nvPr/p:ph", NS)
        ptype = ph.get("type", "body") if ph is not None else None
        if ptype in skip_types:
            continue
        paras = [_para_text(p) for p in sp.iter("{%s}p" % NS["a"])]
        paras = [p for p in paras if p]
        if paras:
            yield ptype, paras


def pptx_slides_zip(path):
    """Fallback reader: zipfile + XML. Returns list of (title, paragraphs, notes_paragraphs)."""
    out = []
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        slide_files = sorted(
            (n for n in names if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)),
            key=lambda n: int(re.search(r"(\d+)\.xml$", n).group(1)))
        for sf in slide_files:
            root = ET.fromstring(z.read(sf))
            title, paras = None, []
            for ptype, ps in _shape_paragraphs(root):
                if ptype in ("title", "ctrTitle") and title is None:
                    title = " ".join(ps)
                else:
                    paras.extend(ps)
            notes = []
            rels = "ppt/slides/_rels/%s.rels" % os.path.basename(sf)
            if rels in names:
                rroot = ET.fromstring(z.read(rels))
                for rel in rroot:
                    if rel.get("Type", "").endswith("/notesSlide"):
                        target = os.path.normpath(os.path.join("ppt/slides", rel.get("Target"))).replace("\\", "/")
                        if target in names:
                            nroot = ET.fromstring(z.read(target))
                            for ptype, ps in _shape_paragraphs(nroot):
                                notes.extend(ps)
            out.append((title, paras, notes))
    return out


def pptx_slides_pythonpptx(path):
    from pptx import Presentation  # noqa: optional dependency
    prs = Presentation(path)
    out = []
    for slide in prs.slides:
        title, paras = None, []
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            ps = [p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()]
            if not ps:
                continue
            if shape == slide.shapes.title and title is None:
                title = " ".join(ps)
            else:
                paras.extend(ps)
        notes = []
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = [p.text.strip() for p in slide.notes_slide.notes_text_frame.paragraphs if p.text.strip()]
        out.append((title, paras, notes))
    return out, "python-pptx"


def slides_from_pptx(path, title=None):
    reader = "zipfile-xml"
    try:
        raw, reader = pptx_slides_pythonpptx(path)
    except ImportError:
        raw = pptx_slides_zip(path)
    slides = []
    for i, (t, paras, notes) in enumerate(raw):
        heading = t or (paras[0] if paras else "%d" % (i + 1))
        body_paras = paras if t else paras[1:]
        bullets = [short_bullet(p) for p in body_paras[:MAX_BULLETS]]
        extra = body_paras[MAX_BULLETS:]
        full_notes = list(notes) + [p for p in body_paras[:MAX_BULLETS] if len(p) > BULLET_MAX] + extra
        slides.append({"heading": strip_terminal(cut_at_clause(heading, TITLE_MAX, False)), "bullets": bullets,
                       "notes": full_notes, "title": i == 0, "body": None})
    if not slides:
        sys.exit("make-deck: no slides found in %s" % path)
    if title:
        slides[0]["heading"] = title
    return slides[0]["heading"], slides, reader


# ---------------------------------------------------------------- markdown out

def front_matter(title, theme, lang):
    return "\n".join([
        "---",
        "marp: true",
        "theme: %s" % theme,
        "paginate: true",
        "size: 16:9",
        "lang: %s" % lang,
        "title: %s" % json.dumps(title, ensure_ascii=False),
        "---",
        "",
    ])


def slides_to_markdown(title, slides, theme, lang):
    parts = []
    for i, s in enumerate(slides):
        block = []
        if s["title"]:
            block.append("<!-- _class: lead -->")
            block.append("<!-- _paginate: false -->")
            block.append("# %s" % s["heading"])
        else:
            block.append("## %s" % s["heading"])
        if s["body"]:
            block.append("")
            block.append(s["body"])
        elif s["bullets"]:
            block.append("")
            block.extend("- %s" % b for b in s["bullets"])
        if s["notes"]:
            block.append("")
            block.append("<!--\n%s\n-->" % "\n\n".join(n.replace("-->", "- ->") for n in s["notes"]))
        parts.append("\n".join(block))
    return front_matter(title, theme, lang) + "\n" + "\n\n---\n\n".join(parts) + "\n"


def parse_marp_markdown(text):
    """From a Marp .md: (title, slide_count, [(heading, [notes])]) for notes.md and the manifest."""
    fm = {}
    body = text
    m = re.match(r"^---\n(.*?)\n---\n", text, re.S)
    if m:
        body = text[m.end():]
        for ln in m.group(1).split("\n"):
            if ":" in ln:
                k, v = ln.split(":", 1)
                fm[k.strip()] = v.strip().strip('"')
    chunks = re.split(r"(?m)^---\s*$", body)
    info = []
    for ch in chunks:
        if not ch.strip():
            continue
        h = re.search(r"(?m)^#{1,3}\s+(.*)$", ch)
        notes = [n.strip() for n in re.findall(r"<!--(.*?)-->", ch, re.S)]
        notes = [n for n in notes if n and not n.startswith("_")]
        info.append((h.group(1).strip() if h else "", notes))
    title = fm.get("title") or (info[0][0] if info else "deck")
    return title, info


def write_notes_md(path, title, info, lang):
    w = WORDS.get(lang, WORDS["en"])
    lines = ["# %s: %s" % (w["notes"], title), "", "created: %s" % dt.date.today().isoformat(), ""]
    for i, (heading, notes) in enumerate(info, 1):
        lines.append("## %s %d. %s" % (w["slide"], i, heading))
        lines.append("")
        if notes:
            lines.extend(notes)
        else:
            lines.append("—")
        lines.append("")
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))


# ---------------------------------------------------------------- rendering

def marp_env():
    env = dict(os.environ)
    chrome = env.get("CHROME_PATH")
    if not chrome:
        found = sorted(glob.glob(CHROME_GLOB))
        chrome = found[-1] if found else None
    if chrome:
        env["CHROME_PATH"] = chrome
        libs = env.get("CHROME_LIBS") or ":".join(d for d in CHROME_LIBS_DEFAULT.split(":") if os.path.isdir(d))
        if libs and "LD_LIBRARY_PATH" not in env:
            env["LD_LIBRARY_PATH"] = libs
    env.setdefault("CHROME_NO_SANDBOX", "1")  # the box has no usable chromium sandbox (userns off)
    return env


def run_marp(args, cwd, env):
    cmd = os.environ.get("MARP_CMD", MARP_DEFAULT).split() + args
    t0 = time.time()
    r = subprocess.run(cmd, cwd=cwd, env=env, capture_output=True, text=True, stdin=subprocess.DEVNULL)  # marp reads a piped stdin
    secs = round(time.time() - t0, 2)
    if r.returncode != 0:
        sys.stderr.write(r.stdout + r.stderr)
        sys.exit("make-deck: marp failed: %s" % " ".join(args))
    return secs


def marp_version(env):
    try:
        r = subprocess.run(os.environ.get("MARP_CMD", MARP_DEFAULT).split() + ["--version"],
                           env=env, capture_output=True, text=True, timeout=120, stdin=subprocess.DEVNULL)
        return r.stdout.strip() or r.stderr.strip()
    except Exception as e:  # noqa
        return "unknown (%s)" % e


def pdf_fonts(path):
    """Font names embedded in the PDF (BaseFont entries, also inside deflated object streams)."""
    data = open(path, "rb").read()
    found = set(re.findall(rb"/BaseFont\s*/([A-Za-z0-9+\-_,]+)", data))
    for m in re.finditer(rb"stream\r?\n(.*?)endstream", data, re.S):
        try:
            found |= set(re.findall(rb"/BaseFont\s*/([A-Za-z0-9+\-_,]+)", zlib.decompress(m.group(1))))
        except Exception:
            continue
    names = sorted({re.sub(r"^[A-Z]{6}\+", "", f.decode("latin-1")) for f in found})
    return names


def cyrillic_fonts_present():
    try:
        r = subprocess.run(["fc-list"], capture_output=True, text=True, timeout=30)
    except Exception:
        return None
    hits = sorted({ln.split(":")[1].strip().split(",")[0] for ln in r.stdout.splitlines()
                   if re.search(r"dejavu|noto|liberation", ln, re.I) and ":" in ln})
    return hits


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 16), b""):
            h.update(chunk)
    return h.hexdigest()


# ---------------------------------------------------------------- speech in

def _multipart(fields, fname, data, mime):
    b = "----deck" + uuid.uuid4().hex
    out = []
    for k, v in fields.items():
        out.append(('--%s\r\nContent-Disposition: form-data; name="%s"\r\n\r\n%s\r\n' % (b, k, v)).encode())
    out.append(('--%s\r\nContent-Disposition: form-data; name="file"; filename="%s"\r\nContent-Type: %s\r\n\r\n'
                % (b, fname, mime)).encode() + data + b"\r\n")
    out.append(("--%s--\r\n" % b).encode())
    return b"".join(out), "multipart/form-data; boundary=" + b


def transcribe(path):
    """The words of an audio file: CC_STT_URL (OpenAI transcription shape, model CC_STT_MODEL) when set,
    else ElevenLabs speech-to-text with ELEVENLABS_API_KEY. Same backends as the bridge's voice."""
    data = open(path, "rb").read()
    mime = {".mp3": "audio/mpeg", ".m4a": "audio/mp4", ".wav": "audio/wav", ".webm": "audio/webm"}.get(
        os.path.splitext(path)[1].lower(), "audio/ogg")
    url = os.environ.get("CC_STT_URL", "").strip()
    key = os.environ.get("ELEVENLABS_API_KEY", "").strip()
    if url:
        body, ctype = _multipart({"model": os.environ.get("CC_STT_MODEL", "Systran/faster-whisper-small")},
                                 os.path.basename(path), data, mime)
        req = urllib.request.Request(url, data=body, method="POST", headers={"Content-Type": ctype})
    elif key:
        body, ctype = _multipart({"model_id": os.environ.get("CC_ELEVEN_STT_MODEL", "scribe_v1")},
                                 os.path.basename(path), data, mime)
        req = urllib.request.Request("https://api.elevenlabs.io/v1/speech-to-text", data=body, method="POST",
                                     headers={"xi-api-key": key, "Content-Type": ctype})
    else:
        sys.exit("make-deck: audio input needs CC_STT_URL or ELEVENLABS_API_KEY (or transcribe it first)")
    with urllib.request.urlopen(req, timeout=600) as r:
        text = str(json.loads(r.read() or b"{}").get("text") or "").strip()
    if not text:
        sys.exit("make-deck: the transcription came back empty")
    return text


# ---------------------------------------------------------------- editable pptx

def marp_slides(md):
    """deck.md -> [{"heading", "level", "items": [(depth, text)], "notes": [str], "lead": bool}] for the PPTX.
    Headings #/##/###, list items (- * 1.) with their indent depth, other text lines as plain items,
    images and code fences skipped (named in the notes), HTML comments as notes unless they are directives."""
    body = re.sub(r"^---\n.*?\n---\n", "", md, count=1, flags=re.S)
    out = []
    for ch in re.split(r"(?m)^---\s*$", body):
        if not ch.strip():
            continue
        lead = bool(re.search(r"<!--\s*_class:\s*[^>]*lead", ch))
        notes = []
        for c in re.findall(r"<!--(.*?)-->", ch, re.S):
            c = c.strip()
            if c and not re.match(r"^_?[a-zA-Z]+\s*:", c):
                notes.append(c)
        text = re.sub(r"<!--.*?-->", "", ch, flags=re.S)
        text = re.sub(r"(?ms)^```.*?^```", "", text)
        heading, level, items = "", 0, []
        for ln in text.split("\n"):
            if not ln.strip():
                continue
            m = re.match(r"^(#{1,6})\s+(.*)$", ln)
            if m and not heading:
                heading, level = m.group(2).strip(), len(m.group(1))
                continue
            if re.match(r"^\s*!\[", ln):
                notes.append("[image: %s]" % ln.strip())
                continue
            m = re.match(r"^(\s*)(?:[-*+]|\d+[.)])\s+(.*)$", ln)
            if m:
                items.append((min(len(m.group(1).replace("\t", "  ")) // 2, 4), m.group(2).strip()))
            elif m := re.match(r"^#{1,6}\s+(.*)$", ln):
                items.append((0, m.group(1).strip()))
            else:
                items.append((0, ln.strip()))
        items = [(d, re.sub(r"\*\*(.+?)\*\*|\*(.+?)\*|`(.+?)`", lambda m: next(g for g in m.groups() if g), t))
                 for d, t in items]
        out.append({"heading": heading, "level": level, "items": items, "notes": notes, "lead": lead})
    return out


def write_pptx(md, path):
    """An editable 16:9 PPTX from deck.md: title slide for the lead slide, Title and Content otherwise,
    bullets in the body placeholder at their depth, speaker notes in the notes pane. Returns False when
    python-pptx is missing (the caller falls back to Marp's image PPTX)."""
    try:
        from pptx import Presentation
        from pptx.util import Emu, Pt
    except ImportError:
        return False
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    for i, s in enumerate(marp_slides(md)):
        if s["lead"] or (i == 0 and s["level"] == 1):
            sl = prs.slides.add_slide(prs.slide_layouts[0])
            sl.shapes.title.text = s["heading"]
            sub = sl.placeholders[1]
            sub.text = "\n".join(t for _, t in s["items"])
            if not s["items"]:
                sub._element.getparent().remove(sub._element)
        else:
            sl = prs.slides.add_slide(prs.slide_layouts[1] if s["items"] else prs.slide_layouts[5])
            sl.shapes.title.text = s["heading"]
            for ph in sl.placeholders:
                if ph.placeholder_format.idx == 1:
                    tf = ph.text_frame
                    for j, (depth, t) in enumerate(s["items"]):
                        para = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                        para.text, para.level = t, depth
                        if len(s["items"]) > 6:
                            para.font.size = Pt(20)
        for ph in sl.placeholders:  # the default template is 4:3; stretch its boxes across 16:9
            ph.left, ph.width = Emu(int(ph.left * 4 / 3)), Emu(int(ph.width * 4 / 3))
        if s["notes"]:
            sl.notes_slide.notes_text_frame.text = "\n\n".join(s["notes"])
    prs.save(path)
    return True


# ---------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("input")
    ap.add_argument("--out", required=True, help="output root; the deck lands in <out>/<slug>/")
    ap.add_argument("--lang", default="auto", choices=LANGS, help="auto: ru when the text is Cyrillic, else en")
    ap.add_argument("--theme", default="default", choices=THEMES)
    ap.add_argument("--title", default=None)
    ap.add_argument("--slug", default=None, help="folder name under --out (default: from the title)")
    ap.add_argument("--draft-only", action="store_true", help="write deck.md and notes.md, do not render")
    a = ap.parse_args()

    src = os.path.abspath(a.input)
    if not os.path.isfile(src):
        sys.exit("make-deck: no such file: %s" % src)
    ext = os.path.splitext(src)[1].lower()
    t_start = time.time()
    reader = None
    passthrough = False

    transcript = None
    if ext in AUDIO_EXT:
        transcript = transcribe(src)
    if ext == ".pptx":
        title, slides, reader = slides_from_pptx(src, a.title)
        kind = "pptx"
        if a.lang == "auto":
            a.lang = "ru" if mostly_cyrillic(json.dumps(slides, ensure_ascii=False)) else "en"
        md = slides_to_markdown(title, slides, a.theme, a.lang)
    else:
        text = transcript if transcript is not None else open(src, encoding="utf-8", errors="replace").read()
        if a.lang == "auto":
            a.lang = "ru" if mostly_cyrillic(text) else "en"
        if ext == ".md" and re.match(r"^---\n.*?\bmarp:\s*true\b.*?\n---\n", text, re.S):
            kind = "marp-md"
            passthrough = True
            md = text
            if a.title:
                md = re.sub(r"(?m)^title:.*$", "title: %s" % json.dumps(a.title, ensure_ascii=False), md, count=1)
            if not re.search(r"(?m)^theme:", md):
                md = md.replace("marp: true", "marp: true\ntheme: %s" % a.theme, 1)
            title, _ = parse_marp_markdown(md)
        elif ext == ".md" and re.search(r"(?m)^#{1,2}\s", text):
            kind = "plain-md"
            title, slides = slides_from_plain_markdown(text, a.title)
            md = slides_to_markdown(title, slides, a.theme, a.lang)
        else:
            kind = "audio" if transcript is not None else "transcript"
            title, slides = slides_from_transcript(text, a.title, a.lang)
            md = slides_to_markdown(title, slides, a.theme, a.lang)

    slug = a.slug or slugify(title)
    outdir = os.path.join(os.path.abspath(a.out), slug)
    os.makedirs(outdir, exist_ok=True)
    deck_md = os.path.join(outdir, "deck.md")
    if passthrough and os.path.abspath(src) == deck_md:
        pass  # re-rendering the draft in place
    else:
        with open(deck_md, "w", encoding="utf-8") as f:
            f.write(md)
    if transcript is not None:
        with open(os.path.join(outdir, "transcript.txt"), "w", encoding="utf-8") as f:
            f.write(transcript + "\n")
    _, info = parse_marp_markdown(md)
    write_notes_md(os.path.join(outdir, "notes.md"), title, info, a.lang)

    manifest = {
        "created": dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "tool": "scripts/deck/make-deck.py v1.1",
        "input": {"path": src, "kind": kind, "bytes": os.path.getsize(src), "sha256": sha256(src),
                  "pptx_reader": reader},
        "title": title, "slug": slug, "lang": a.lang,
        "theme": a.theme if not passthrough else (re.search(r"(?m)^theme:\s*(\S+)", md) or [None, "?"])[1],
        "slide_count": len(info),
        "cyrillic": has_cyrillic(md),
        "stage": "draft" if a.draft_only else "rendered",
        "outputs": {"deck_md": deck_md, "notes_md": os.path.join(outdir, "notes.md")},
        "render_seconds": {},
    }

    if a.draft_only:
        manifest["seconds_total"] = round(time.time() - t_start, 2)
        with open(os.path.join(outdir, "manifest.json"), "w", encoding="utf-8") as f:
            json.dump(manifest, f, ensure_ascii=False, indent=2)
        print("draft: %s (%d slides). Edit it, then: make-deck.py %s --out %s --slug %s"
              % (deck_md, len(info), deck_md, os.path.abspath(a.out), slug))
        return

    env = marp_env()
    preview = os.path.join(outdir, "preview")
    if os.path.isdir(preview):
        shutil.rmtree(preview)
    os.makedirs(preview)
    manifest["marp_version"] = marp_version(env)
    manifest["render_seconds"]["pdf"] = run_marp(["deck.md", "--pdf", "--pdf-notes", "-o", "deck.pdf"], outdir, env)
    t0 = time.time()
    if write_pptx(md, os.path.join(outdir, "deck.pptx")):
        manifest["pptx_kind"] = "editable (python-pptx: title and bullet placeholders, speaker notes)"
    else:
        run_marp(["deck.md", "--pptx", "-o", "deck.pptx"], outdir, env)
        manifest["pptx_kind"] = ("images (Marp: one picture per slide, text not editable; install python-pptx "
                                 "in the deck venv for an editable PPTX, docs/DECK.md)")
    manifest["render_seconds"]["pptx"] = round(time.time() - t0, 2)
    manifest["render_seconds"]["png"] = run_marp(["deck.md", "--images", "png", "-o", "preview/deck.png"],
                                                 outdir, env)
    pngs = sorted(p for p in os.listdir(preview) if p.endswith(".png"))
    manifest["outputs"].update({
        "deck_pdf": os.path.join(outdir, "deck.pdf"), "deck_pptx": os.path.join(outdir, "deck.pptx"),
        "preview": [os.path.join(preview, p) for p in pngs],
    })
    manifest["fonts"] = {"pdf_embedded": pdf_fonts(os.path.join(outdir, "deck.pdf")),
                         "cyrillic_capable_installed": cyrillic_fonts_present()}
    manifest["seconds_total"] = round(time.time() - t_start, 2)
    if manifest["cyrillic"] and not manifest["fonts"]["pdf_embedded"]:
        manifest["warning"] = "Cyrillic text but no embedded font found in the PDF; check fc-list."
    if len(pngs) != manifest["slide_count"]:
        manifest["warning"] = "preview has %d images but %d slides were counted" % (len(pngs), manifest["slide_count"])
    with open(os.path.join(outdir, "manifest.json"), "w", encoding="utf-8") as f:
        json.dump(manifest, f, ensure_ascii=False, indent=2)
    print("deck: %s  slides=%d  pptx=%s  pdf=%ss pptx=%ss png=%ss  fonts=%s%s" % (
        outdir, manifest["slide_count"], manifest["pptx_kind"].split(" ")[0],
        manifest["render_seconds"]["pdf"], manifest["render_seconds"]["pptx"],
        manifest["render_seconds"]["png"], ",".join(manifest["fonts"]["pdf_embedded"]) or "none",
        ("  WARNING: " + manifest["warning"]) if "warning" in manifest else ""))


if __name__ == "__main__":
    main()
